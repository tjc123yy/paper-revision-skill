# -*- coding: utf-8 -*-
"""rewrite_tables.py —— 按配置重写 docx 里的表格内容与结构，用于降低查重重复率。

背景（2026-10-07 学生D任务）：导师反馈"表4-1到表4-4、表6-1、表6-2 的查重重复率太高"。
数据字典表（列名/数据类型/长度/是否主键/是否可空/说明）是全网模板，逐字高度雷同；
测试用例表的措辞也是通用模板。只改个别词没用，必须**同时**换掉：列结构、表头措辞、
字段名与含义、说明文字、测试用例的描述与结果。

本脚本把那次的做法参数化，任何论文都可复用。它做四件事：

1. **替换单元格文字**：逐单元格写入新文本，段落沿用原单元格模板（pStyle 42 +
   宋体/Times New Roman 五号 + 行距 360 auto），所以版式不会漂；
2. **改列结构**：`drop_cols` 删列、`add_cols` 插列（列宽按 `widths` 重排，
   并同步 tblGrid 与每个 tcW）；
3. **改行数**：按 `rows` 的条数自动补/裁行（克隆首行外壳，清掉重复的 w14:paraId）；
4. **改成三线表并拆纵向合并**：`three_line=true` 去网格线（顶线 12 / 表头下细线 6 /
   底线 12，表头加底纹与 tblHeader）；`flatten_vmerge` 把首列 vMerge 拆成普通单元格，
   避免跨页时合并标签错位（LibreOffice 实测会把组名挪到分页点，出现"第一行空、
   第二行显示上一组名"）。

配置格式（JSON，UTF-8）::

    {
      "tables": [
        {
          "index": 0,                       // body 内表格序号（0 起）
          "rows": [["字段名","含义",...], ["AutoID","景点记录自增编号",...], ...],
          "widths": [1500,1500,1420,1430,1954],   // twips，长度 = 目标列数
          "drop_cols": [4],                 // 先删列
          "add_cols": [],                   // 再插列，元素为 {"at": 3, "width": 1404}
          "left_cols": [1,4],               // 左对齐的列号，其余居中
          "three_line": true,
          "flatten_vmerge": false
        }
      ]
    }

`rows` 的列数必须等于"原列数 - len(drop_cols) + len(add_cols)"，否则脚本直接报错退出。

用法::

    python rewrite_tables.py 原文件.docx 产物.docx --config table_spec.json

自检（脚本内，失败即非 0 退出）：
  * 结构不变量 `drawing/pict/bookmarkStart/bookmarkEnd/fldChar/instrText/hyperlink/sectPr/tbl`
    以及**正文级**段落数必须与源文件完全一致（表格内段落数会随列结构合法变化）；
  * 正文（表格外）段落逐段文字必须一字不变；
  * OOXML 子元素顺序（tblPr/trPr/tcPr/pPr）必须合法，否则 Word 会提示"需修复"；
  * 无嵌套 `<w:r>`。
"""
import argparse
import copy
import json
import sys

try:
    from lxml import etree
except ImportError:
    raise SystemExit('需要 lxml：请用 paper-revision/scripts/run.ps1 执行本脚本')

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def q(tag):
    return '{%s}%s' % (W, tag)


TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd',
              'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter',
              'cantSplit', 'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden']
PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
             'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
             'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
             'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
             'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection',
             'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
             'sectPr', 'pPrChange']


def inserter(parent, child, order):
    """按 schema 顺序插入 child（顺序错会被 Word 判为需修复）。"""
    idx = order.index(child.tag.split('}')[-1])
    for i, el in enumerate(parent):
        name = el.tag.split('}')[-1]
        if name in order and order.index(name) > idx:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def import_docxlib(path):
    import os
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import docxlib
    return docxlib


D = import_docxlib(__file__)
from docx import Document  # noqa: E402


def clear_paraid(el):
    for sub in el.iter():
        for a in list(sub.attrib):
            if a.endswith('}paraId') or a.endswith('}textId'):
                del sub.attrib[a]


def set_border(parent, edge, val, sz=None):
    el = parent.find(q(edge))
    if el is None:
        el = etree.SubElement(parent, q(edge))
    el.set(q('val'), val)
    if val in ('nil', 'none'):
        for a in (q('sz'), q('color')):
            if el.get(a) is not None:
                del el.attrib[a]
    else:
        el.set(q('sz'), str(sz))
        el.set(q('color'), '000000')
    return el


def para_template(tbl):
    """首个单元格段落作模板（保留 pStyle 与 rPr），清空文字 run。"""
    p = tbl.find(q('tr')).find(q('tc')).find(q('p'))
    tmpl = copy.deepcopy(p)
    for r in tmpl.findall(q('r')):
        tmpl.remove(r)
    for h in tmpl.findall(q('hyperlink')):
        tmpl.remove(h)
    return tmpl


def set_cell(tc, text, tmpl, rpr, align):
    for child in list(tc):
        if child.tag != q('tcPr'):
            tc.remove(child)
    p = copy.deepcopy(tmpl)
    pPr = p.find(q('pPr'))
    if pPr is None:
        pPr = etree.SubElement(p, q('pPr'))
    jc = pPr.find(q('jc'))
    if jc is None:
        jc = etree.Element(q('jc'))
        inserter(pPr, jc, PPR_ORDER)
    jc.set(q('val'), align)
    for r in p.findall(q('r')):
        p.remove(r)
    D.append_run(p, text, rpr)
    tc.append(p)


def drop_col(tbl, idx):
    grid = tbl.find(q('tblGrid'))
    grid.remove(grid.findall(q('gridCol'))[idx])
    for tr in tbl.findall(q('tr')):
        tcs = tr.findall(q('tc'))
        tr.remove(tcs[idx])


def add_col(tbl, idx, width):
    grid = tbl.find(q('tblGrid'))
    cols = grid.findall(q('gridCol'))
    newcol = copy.deepcopy(cols[0])
    newcol.set(q('w'), str(width))
    cols[idx - 1].addnext(newcol)
    for tr in tbl.findall(q('tr')):
        tcs = tr.findall(q('tc'))
        newtc = copy.deepcopy(tcs[0])
        clear_paraid(newtc)
        tcs[idx - 1].addnext(newtc)


def set_widths(tbl, widths):
    grid = tbl.find(q('tblGrid'))
    cols = grid.findall(q('gridCol'))
    if len(cols) != len(widths):
        raise SystemExit('gridCol %d != widths %d' % (len(cols), len(widths)))
    for c, w in zip(cols, widths):
        c.set(q('w'), str(w))
    for tr in tbl.findall(q('tr')):
        for tc, w in zip(tr.findall(q('tc')), widths):
            tcPr = tc.find(q('tcPr'))
            cw = tcPr.find(q('tcW'))
            if cw is None:
                cw = etree.Element(q('tcW'))
                inserter(tcPr, cw, TCPR_ORDER)
            cw.set(q('w'), str(w))
            cw.set(q('type'), 'dxa')
    tw = tbl.find(q('tblPr')).find(q('tblW'))
    if tw is not None:
        tw.set(q('w'), str(sum(widths)))
        tw.set(q('type'), 'dxa')


def three_line(tbl):
    tblPr = tbl.find(q('tblPr'))
    for ex in tbl.findall(q('tblPrEx')):
        tblPr.remove(ex)
    borders = tblPr.find(q('tblBorders'))
    if borders is None:
        borders = etree.SubElement(tblPr, q('tblBorders'))
    for edge in ('top', 'bottom'):
        set_border(borders, edge, 'single', 12)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        set_border(borders, edge, 'none', 0)
    for tr in tbl.findall(q('tr')):
        ex = tr.find(q('tblPrEx'))
        if ex is not None:                       # 行级边框会盖过表级设置
            tr.remove(ex)
        for tc in tr.findall(q('tc')):
            tcPr = tc.find(q('tcPr'))
            if tcPr is None:
                continue
            tcb = tcPr.find(q('tcBorders'))
            if tcb is not None:
                tcPr.remove(tcb)
    trs = tbl.findall(q('tr'))
    if not trs:
        return
    trPr = trs[0].find(q('trPr'))
    if trPr is None:
        trPr = etree.Element(q('trPr'))
        trs[0].insert(0, trPr)
    if trPr.find(q('tblHeader')) is None:
        inserter(trPr, etree.Element(q('tblHeader')), TRPR_ORDER)
    for tc in trs[0].findall(q('tc')):
        tcPr = tc.find(q('tcPr'))
        if tcPr is None:
            tcPr = etree.Element(q('tcPr'))
            tc.insert(0, tcPr)
        tcb = etree.Element(q('tcBorders'))
        set_border(tcb, 'bottom', 'single', 6)
        inserter(tcPr, tcb, TCPR_ORDER)
        shd = tcPr.find(q('shd'))
        if shd is None:
            shd = etree.Element(q('shd'))
            shd.set(q('val'), 'clear')
            shd.set(q('color'), 'auto')
            inserter(tcPr, shd, TCPR_ORDER)
        shd.set(q('fill'), 'F2F2F2')


def no_split_rows(tbl):
    """数据行禁止跨页断开，避免"表头+半行"孤零零留在上一页。"""
    for tr in tbl.findall(q('tr')):
        trPr = tr.find(q('trPr'))
        if trPr is None:
            trPr = etree.Element(q('trPr'))
            tr.insert(0, trPr)
        if trPr.find(q('cantSplit')) is None:
            inserter(trPr, etree.Element(q('cantSplit')), TRPR_ORDER)


def flatten_vmerge(tbl, col=0):
    """把首列纵向合并拆成普通单元格（跨页时合并标签会错位）。"""
    n = 0
    for tr in tbl.findall(q('tr')):
        tcs = tr.findall(q('tc'))
        if col >= len(tcs):
            continue
        tcPr = tcs[col].find(q('tcPr'))
        if tcPr is None:
            continue
        vm = tcPr.find(q('vMerge'))
        if vm is not None:
            tcPr.remove(vm)
            n += 1
    return n


def set_rows(tbl, rows_data, tmpl, rpr, left_cols):
    trs = tbl.findall(q('tr'))
    base_tr = trs[0]
    while len(tbl.findall(q('tr'))) > len(rows_data):
        tbl.remove(tbl.findall(q('tr'))[-1])
    while len(tbl.findall(q('tr'))) < len(rows_data):
        new_tr = copy.deepcopy(base_tr)
        clear_paraid(new_tr)
        tbl.append(new_tr)
    for ri, (tr, row_data) in enumerate(zip(tbl.findall(q('tr')), rows_data)):
        tcs = tr.findall(q('tc'))
        if len(tcs) != len(row_data):
            raise SystemExit('第 %d 行：单元格 %d 个，数据 %d 列，不匹配' % (ri, len(tcs), len(row_data)))
        for ci, (tc, text) in enumerate(zip(tcs, row_data)):
            tcPr = tc.find(q('tcPr'))
            if tcPr is not None and tcPr.find(q('vAlign')) is None:
                va = etree.Element(q('vAlign'))
                va.set(q('val'), 'center')
                inserter(tcPr, va, TCPR_ORDER)
            set_cell(tc, text, tmpl, rpr, 'left' if ci in left_cols else 'center')


def count_invariants(root):
    keys = ['drawing', 'pict', 'object', 'bookmarkStart', 'bookmarkEnd', 'fldChar',
            'instrText', 'hyperlink', 'sectPr', 'tbl']
    d = {k: len(list(root.iter(q(k)))) for k in keys}
    d['body_p'] = len([c for c in root.find(q('body')) if c.tag == q('p')])
    return d


def body_texts(root):
    return [D.get_text(p) for p in root.find(q('body')) if p.tag == q('p')]


def check_order(root):
    ORDER = {
        'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
                  'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders',
                  'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
        'trPr': TRPR_ORDER,
        'tcPr': TCPR_ORDER,
        'pPr': PPR_ORDER,
        'rPr': ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike',
                'dstrike', 'outline', 'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid',
                'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs',
                'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs',
                'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath'],
    }
    bad = []
    for el in root.iter():
        name = el.tag.split('}')[-1]
        if name not in ORDER:
            continue
        seq = [c.tag.split('}')[-1] for c in el]
        known = [x for x in seq if x in ORDER[name]]
        idx = [ORDER[name].index(x) for x in known]
        if idx != sorted(idx):
            bad.append((name, seq))
    return bad


def main():
    ap = argparse.ArgumentParser(description='按配置重写 docx 表格内容与结构')
    ap.add_argument('src')
    ap.add_argument('dst')
    ap.add_argument('--config', required=True)
    ap.add_argument('--no-cant-split', action='store_true', help='不给数据行加 cantSplit')
    ap.add_argument('--skip-three-line', action='store_true', help='保留原网格线')
    a = ap.parse_args()

    spec = json.load(open(a.config, encoding='utf-8'))
    src_doc = Document(a.src)
    doc = Document(a.src)
    if len(doc.tables) != len(src_doc.tables):
        raise SystemExit('表格数异常')
    base_inv = count_invariants(src_doc.element)
    base_body = body_texts(src_doc.element)

    for item in spec['tables']:
        ti = item['index']
        tbl = doc.tables[ti]._tbl
        tmpl = para_template(tbl)
        rpr = tmpl.find(q('pPr')).find(q('rPr'))
        if rpr is None:
            raise SystemExit('表 %d 模板缺 rPr，无法保持字体' % ti)
        rpr = copy.deepcopy(rpr)

        n_cols = len(tbl.findall(q('tr'))[0].findall(q('tc')))
        for c in sorted(item.get('drop_cols', []), reverse=True):
            drop_col(tbl, c)
            n_cols -= 1
        for c in item.get('add_cols', []):
            add_col(tbl, c['at'], c['width'])
            n_cols += 1
        rows = item['rows']
        if len(rows[0]) != n_cols:
            raise SystemExit('表 %d：数据 %d 列，实际结构 %d 列' % (ti, len(rows[0]), n_cols))
        if item.get('widths'):
            if len(item['widths']) != n_cols:
                raise SystemExit('表 %d：widths %d 个 != %d 列' % (ti, len(item['widths']), n_cols))
            set_widths(tbl, item['widths'])
        set_rows(tbl, rows, tmpl, rpr, set(item.get('left_cols', [])))
        if item.get('flatten_vmerge'):
            flatten_vmerge(tbl, 0)
        if item.get('three_line', True) and not a.skip_three_line:
            three_line(tbl)
        if not a.no_cant_split:
            no_split_rows(tbl)
        print('表 %d -> %d 行 x %d 列 %s' % (ti, len(rows), len(rows[0]),
                                             '(三线表)' if item.get('three_line', True) else ''))

    doc.save(a.dst)

    # ---- 自检 ----
    out_doc = Document(a.dst)
    problems = []
    inv = count_invariants(out_doc.element)
    for k, v in base_inv.items():
        if inv[k] != v:
            problems.append('结构不变量 %s: %s -> %s' % (k, v, inv[k]))
    if body_texts(out_doc.element) != base_body:
        problems.append('正文（表格外）段落文字发生变化')
    nested = D.check_nested_runs(out_doc.element)
    if nested:
        problems.append('存在嵌套 <w:r>: %d' % nested)
    bad_order = check_order(out_doc.element)
    if bad_order:
        problems.append('XML 子元素顺序非法: %s' % bad_order[:3])
    if problems:
        for p in problems:
            print('FAIL:', p)
        raise SystemExit(1)
    print('SELF-CHECK PASS ->', a.dst)


if __name__ == '__main__':
    main()
